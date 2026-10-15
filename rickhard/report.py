"""Formatted Excel report (openpyxl): one workbook with everything a team or a
manager needs, readable without the app."""

from __future__ import annotations

import io
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import APP_NAME, __version__
from .constants import DEV_COLS

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="1F3A5F")
SUB_FONT = Font(size=10, color="666666")
GOOD_FILL = PatternFill("solid", fgColor="D9F2D9")
BAD_FILL = PatternFill("solid", fgColor="F8D7D7")
WARN_FILL = PatternFill("solid", fgColor="FFF1CC")
GREY_FILL = PatternFill("solid", fgColor="EDEDED")
THIN = Border(bottom=Side(style="thin", color="DDDDDD"))


def _clean(df: pd.DataFrame) -> pd.DataFrame:
    d = df.copy()
    for c in d.columns:
        if pd.api.types.is_datetime64_any_dtype(d[c]):
            if getattr(d[c].dt, "tz", None) is not None:
                d[c] = d[c].dt.tz_localize(None)
        elif d[c].dtype == object:
            d[c] = d[c].map(lambda v: v if v is None or isinstance(v, (str, int, float, bool, pd.Timestamp, datetime))
                            else str(v))
    return d.replace({np.inf: np.nan, -np.inf: np.nan})


def _write_table(ws, df: pd.DataFrame, start_row: int = 1, limit: float | None = None, number_format: str = "0.00"):
    df = _clean(df)
    cols = list(df.columns)
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=start_row, column=j, value=str(c))
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, v in enumerate(row, start=1):
            if isinstance(v, float) and np.isnan(v):
                v = None
            elif v is pd.NaT:
                v = None
            cell = ws.cell(row=i, column=j, value=v)
            if isinstance(v, float):
                cell.number_format = number_format
            elif isinstance(v, (pd.Timestamp, datetime)):
                cell.number_format = "yyyy-mm-dd hh:mm"
    n = len(df)
    last = start_row + max(n, 1)
    for j, c in enumerate(cols, start=1):
        width = max([len(str(c))] + [len(str(x)) for x in df[c].head(200).tolist()]) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(48, max(9, width))
    if n:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{last}"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    # conditional formats
    for j, c in enumerate(cols, start=1):
        col = get_column_letter(j)
        rng = f"{col}{start_row + 1}:{col}{last}"
        if limit is not None and (c in DEV_COLS or (str(c).endswith("_Sim") and str(c)[:-4] in DEV_COLS)):
            ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=[str(limit)], fill=BAD_FILL))
            ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=[str(-limit)], fill=BAD_FILL))
        if str(c).startswith("Status") or c in ("SquarenessStatus", "State", "Verdict"):
            first = f"{col}{start_row + 1}"
            for word, fill in (("PASS", GOOD_FILL), ("FAIL", BAD_FILL), ("INCOMPLETE", GREY_FILL), ("DEFORMED", BAD_FILL),
                               ("SQUARE OK", GOOD_FILL), ("Signal", BAD_FILL), ("Watch", WARN_FILL),
                               ("In control", GOOD_FILL), ("Not acceptable", BAD_FILL), ("Good", GOOD_FILL)):
                ws.conditional_formatting.add(rng, FormulaRule(formula=[f'EXACT({first},"{word}")'], fill=fill))
    return last


def _sheet(wb, name: str, df: pd.DataFrame | None, note: str | None = None, **kw):
    if df is None or (isinstance(df, pd.DataFrame) and df.empty):
        return
    ws = wb.create_sheet(name[:31])
    start = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = SUB_FONT
        start = 3
    _write_table(ws, df, start_row=start, **kw)


def first_run_quality_summary(first: pd.DataFrame, incomplete_as_fail: bool) -> pd.DataFrame:
    """The original app's 'First_Run_Quality_Summary' table."""
    n = len(first)
    ok = int((first["Status"] == "PASS").sum()) if n else 0
    nok = int((first["Status"] == "FAIL").sum()) if n else 0
    inc = int((first["Status"] == "INCOMPLETE").sum()) if n else 0
    rows = [("Total Unique Modules Evaluated", str(n)), ("Passed First Inspection (OK)", str(ok)),
            ("Failed First Inspection (NOK)", str(nok))]
    if not incomplete_as_fail:
        rows.append(("Incomplete Measurements", str(inc)))
    rows.append(("First Pass Yield (FPY)", f"{ok / n * 100:.1f}%" if n else "n/a"))
    return pd.DataFrame(rows, columns=["Metric", "Value"])


def build_excel(a, insights=None, actions=None, extras: dict | None = None, health: dict | None = None,
                events: pd.DataFrame | None = None, comp_rows: pd.DataFrame | None = None,
                title: str = "ST020 corner deviation report", comp_offsets: pd.DataFrame | None = None,
                comp_sims: pd.DataFrame | None = None) -> bytes:
    """``a`` is a rickhard.analysis.Analysis.

    The sheets of the original app's export come first, under the same names
    (Module_Summary_Report, First_Run_Quality_Summary, Weekly_FPY_Trend, …), so
    anything built on that file keeps working; the deeper analyses follow."""
    ex = extras or {}
    r = a.settings.rules
    k = a.kpis
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    f = a.first
    span = (f"{f['Date'].min():%Y-%m-%d} – {f['Date'].max():%Y-%m-%d}" if len(f) else "–")
    ws["A2"] = (f"{APP_NAME} v{__version__} · generated {datetime.now():%Y-%m-%d %H:%M} · {span} · "
                f"{a.scope.label_types()} · tolerance ±{r.spec_limit:g} mm · FPY target {r.fpy_target:g}%")
    ws["A2"].font = SUB_FONT
    rows = [
        ("First-pass yield (FPY)", k.get("fpy"), "%"),
        ("95% confidence interval", f"{k['fpy_ci'][0]:.1f}% – {k['fpy_ci'][1]:.1f}%", ""),
        ("Gap to target", k.get("gap"), "pp"),
        ("Modules (first run)", k.get("modules"), ""),
        ("Measurement runs", k.get("runs"), ""),
        ("Modules re-measured", k.get("remeasured"), ""),
        ("NOK at first measurement", k.get("fail"), ""),
        ("Incomplete first runs (kept apart)", k.get("incomplete"), ""),
        ("Near misses (pass with < margin)", k.get("near_miss"), f"< {r.near_miss:g} mm"),
        ("Weakest corner", k.get("worst_coord", "–"), f"Ppk {k.get('worst_ppk', float('nan')):.2f}"),
        ("Deformed modules", k.get("deformed_rate"), "%"),
        ("Latest week", k.get("last_week", "–"), f"{k.get('last_week_fpy', float('nan')):.1f}% FPY"),
    ]
    ws.cell(row=4, column=1, value="Key figures").font = Font(bold=True, size=12)
    for i, (lab, val, unit) in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=lab)
        cell = ws.cell(row=i, column=2, value=(round(float(val), 2) if isinstance(val, (float, np.floating)) else val))
        cell.font = Font(bold=True)
        ws.cell(row=i, column=3, value=unit).font = SUB_FONT
        for j in (1, 2, 3):
            ws.cell(row=i, column=j).border = THIN
    row = 5 + len(rows) + 1
    if insights:
        ws.cell(row=row, column=1, value="What the data says").font = Font(bold=True, size=12)
        row += 1
        for ins in insights[:12]:
            fill = {"good": GOOD_FILL, "warn": WARN_FILL, "bad": BAD_FILL}.get(ins.level)
            c1 = ws.cell(row=row, column=1, value=ins.title)
            c1.font = Font(bold=True)
            if fill:
                c1.fill = fill
            c2 = ws.cell(row=row, column=2, value=ins.text.replace("**", ""))
            c2.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
            ws.row_dimensions[row].height = 48
            row += 1
        row += 1
    if actions:
        ws.cell(row=row, column=1, value="Action plan").font = Font(bold=True, size=12)
        row += 1
        for j, h in enumerate(["#", "Action", "Why", "Est. FPY gain (pp)", "Owner"], start=1):
            c = ws.cell(row=row, column=j, value=h)
            c.fill, c.font = HEAD_FILL, HEAD_FONT
        row += 1
        for n, act in enumerate(actions[:12], start=1):
            ws.cell(row=row, column=1, value=n)
            ws.cell(row=row, column=2, value=act.action.replace("**", "")).alignment = Alignment(wrap_text=True)
            ws.cell(row=row, column=3, value=act.why.replace("**", "")).alignment = Alignment(wrap_text=True)
            ws.cell(row=row, column=4, value=round(float(act.impact), 1) if act.impact is not None
                    and np.isfinite(act.impact) else None)
            ws.cell(row=row, column=5, value=act.owner)
            ws.row_dimensions[row].height = 32
            row += 1
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 22

    # ---- the original export's sheets, same names ---------------------------------------------
    run_cols = ["Date", "CalendarWeek", "PartID", "BatteryType", "RunNum", *DEV_COLS, "CornersOutOfSpec",
                "MissingCorners", "IsComplete", "Status"]
    _sheet(wb, "Module_Summary_Report", a.runs[[c for c in run_cols if c in a.runs.columns]], limit=r.spec_limit,
           note="All measurement runs in chronological order. Red cells: outside the tolerance.")
    _sheet(wb, "First_Run_Quality_Summary", first_run_quality_summary(f, r.incomplete_as_fail))
    wk = a.weekly.rename(columns={"PassRate": "FPY %", "MA_FPY": "MA FPY %", "CI_lo": "CI low %", "CI_hi": "CI high %",
                                  "NearMissRate": "Near miss %"})
    _sheet(wb, "Weekly_FPY_Trend", wk.drop(columns=["WeekKey", "LowSample"], errors="ignore"), number_format="0.0")
    _sheet(wb, "Failure_Pareto", a.pareto, number_format="0.0")
    _sheet(wb, "Corner_Capability", a.capability.drop(columns=["Coord"], errors="ignore"))
    sq_cols = ["Date", "CalendarWeek", "PartID", "RunNum", "BatteryType", "Diag1", "Diag2", "DeltaDiag", "WidthDelta",
               "LengthDelta", "AngleDevFL", "SquarenessStatus", "RootCause", "Status", "NokCause", "RigidMax",
               "ShapeMax", "PassIfSquare", "PassIfPlaced"]
    _sheet(wb, "Squareness_Analysis", a.sq_runs[[c for c in sq_cols if c in a.sq_runs.columns]])
    vec_cols = ["Date", "CalendarWeek", "PartID", "BatteryType", "Centroid_X", "Centroid_Y", "Vector_Magnitude",
                "Rotation_Angle", "Fit_Tx", "Fit_Ty", "Fit_Rot", "Status"]
    _sheet(wb, "Vector_Drift_Analysis", f[[c for c in vec_cols if c in f.columns]], number_format="0.000")
    if comp_rows is not None:
        _sheet(wb, "Compensation_Summary", comp_rows, number_format="0.000")
    if comp_offsets is not None:
        _sheet(wb, "Corner_Offsets", comp_offsets, number_format="0.000")
    if comp_sims is not None:
        _sheet(wb, "Simulated_Results", comp_sims, limit=r.spec_limit,
               note="First runs of the compensation window, measured and with the chosen compensation (_Sim).")

    # ---- deeper analyses ---------------------------------------------------------------------
    first_cols = ["Date", "CalendarWeek", "Shift", "PartID", "BatteryType", "RunNum", *DEV_COLS, "WorstCoord",
                  "WorstAbs", "Margin", "Status", "Mechanism", "MechanismDetail", "NokCause", "Fit_Tx", "Fit_Ty",
                  "Fit_Rot", "Centroid_X", "Centroid_Y", "Rotation_Angle", "ShapeMax", "DeltaDiag", "SquarenessStatus",
                  "RootCause"]
    _sheet(wb, "First runs", f[[c for c in first_cols if c in f.columns]], limit=r.spec_limit,
           note="First measurement of every module (the basis of FPY), with why each NOK failed.")
    _sheet(wb, "Failure mechanisms", a.mechanisms, number_format="0.0")
    if ex.get("budget") is not None:
        _sheet(wb, "Error budget", ex["budget"], number_format="0.000")
    if ex.get("change_points") is not None:
        _sheet(wb, "Process shifts", ex["change_points"], number_format="0.000")
    if ex.get("alarms") is not None:
        _sheet(wb, "SPC alarms", ex["alarms"], number_format="0.000")
    if ex.get("repeatability") is not None:
        _sheet(wb, "Repeatability (MSA)", ex["repeatability"], number_format="0.000")
    if events is not None and len(events):
        _sheet(wb, "Interventions", events.drop(columns=["WeekKey"], errors="ignore"))
    if health:
        hrows = [(k_, str(v)) for k_, v in health.items() if not isinstance(v, (list, dict)) or k_ in
                 ("limits_in_file", "types", "time_formats", "similar_part_ids")]
        _sheet(wb, "Data health", pd.DataFrame(hrows, columns=["Item", "Value"]))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
