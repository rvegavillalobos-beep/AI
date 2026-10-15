"""The Excel workbook and the HTML report build and contain what they promise."""

import io
import re

import openpyxl
import pandas as pd
import pytest

from rickhard import demo, events, insights, report, report_html


@pytest.fixture(scope="module")
def bundle(demo_analysis, demo_ds):
    return insights.bundle(demo_analysis, demo_ds.health)


@pytest.fixture(scope="module")
def event_table():
    return events.normalize_events(demo.demo_events())


def test_excel_report(demo_analysis, demo_ds, bundle, event_table):
    comp_rows = pd.DataFrame([{"BatteryType": "Type S", "Method": "Robust optimum", "X [mm]": -0.2,
                               "Y [mm]": 0.1, "Yaw [°]": 0.01}])
    data = report.build_excel(demo_analysis, bundle["insights"], bundle["actions"], bundle["extras"], demo_ds.health,
                              event_table, comp_rows, title="Test report")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    names = wb.sheetnames
    # the original export's sheets come first, under the same names
    original = ["Module_Summary_Report", "First_Run_Quality_Summary", "Weekly_FPY_Trend", "Failure_Pareto",
                "Corner_Capability", "Squareness_Analysis", "Vector_Drift_Analysis", "Compensation_Summary"]
    assert names[1:1 + len(original)] == original
    for sheet in ("First runs", "Failure mechanisms", "Process shifts", "SPC alarms", "Interventions", "Data health"):
        assert sheet in names, sheet
    kpi = {wb["First_Run_Quality_Summary"].cell(row=r_, column=1).value:
           wb["First_Run_Quality_Summary"].cell(row=r_, column=2).value for r_ in range(1, 8)}
    assert kpi["Total Unique Modules Evaluated"] == str(len(demo_analysis.first))
    assert kpi["First Pass Yield (FPY)"] == f"{demo_analysis.fpy:.1f}%"
    ws = wb["First runs"]
    assert ws.max_row - 1 >= len(demo_analysis.first)          # one row per module (+ header / notes)
    assert all(len(n) <= 31 for n in names)


@pytest.mark.parametrize("offline", [False, True])
def test_html_report(demo_analysis, bundle, event_table, offline):
    html = report_html.build_html(demo_analysis, bundle["insights"], bundle["actions"], event_table, None,
                                  title="Line 3 — week review", offline=offline, note="Prepared for the Monday meeting")
    assert html.lstrip().lower().startswith("<!doctype html")
    assert "Line 3 — week review" in html and "Prepared for the Monday meeting" in html
    assert html.count("Plotly.newPlot") >= 4
    assert bool(re.search(r'<script[^>]+src="https://cdn\.plot\.ly/', html)) != offline
    assert (len(html) > 3_000_000) == offline                    # offline embeds plotly.js
    text = re.sub(r"<script.*?</script>", "", html, flags=re.S)
    assert not re.search(r">\s*nan\s*<", text, flags=re.I)        # no stray NaN in visible cells
    assert f"{demo_analysis.fpy:.1f}" in text
